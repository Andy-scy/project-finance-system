"""Excel / CSV 导入导出（openpyxl）。金额列单位一律为「元」。"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..models import Project
from .calc import compute_project, load_context
from .money import parse_date_any, parse_money_text, pct_to_bp, yuan_to_cents
from .settings_svc import get_settings

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="微软雅黑", size=10)


def _style_header(ws, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _autowidth(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fmt(v):
    if v is None:
        return None
    if isinstance(v, date):
        return v
    if isinstance(v, float):
        return round(v, 2)
    return v


# ── 导出 ─────────────────────────────────────────────────────

OVERVIEW_HEADERS = [
    "项目名称", "项目编号", "客户名称", "项目负责人", "项目类型", "状态",
    "项目开始日期", "项目结束日期", "归属年度",
    "合同编号", "合同金额(元)", "税率(%)", "不含税金额(元)", "税额(元)", "合同签订日期",
    "已到账(元)", "待到账(元)", "回款率(%)",
    "直接成本(元)", "间接成本(元)", "税收成本(元)", "总成本(元)", "利润(元)", "利润率(%)",
    "已开票(元)", "未开票(元)", "成本回收日期", "质保金金额(元)",
    "预计质保金到账", "质保金状态", "下一笔预计到账", "备注",
]


def export_projects_xlsx(db: Session, year: int | None = None) -> bytes:
    settings = get_settings(db)
    q = db.query(Project).order_by(Project.id)
    if year:
        q = q.filter(Project.year == year)
    projects = q.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "项目总览"
    _style_header(ws, OVERVIEW_HEADERS)
    _autowidth(ws, [28, 14, 20, 12, 12, 8, 12, 12, 8, 16, 14, 8, 14, 12, 12,
                    14, 14, 10, 14, 14, 14, 14, 10, 14, 14, 12, 14, 12, 12, 14, 24])
    money_cols = {11, 13, 14, 16, 17, 19, 20, 21, 22, 23, 25, 26, 28}
    date_cols = {7, 8, 15, 27, 29, 31}
    r = 2
    for p in projects:
        cp = compute_project(load_context(db, p), settings)
        c = p.contract
        row = [
            p.name, p.code, p.customer_name, p.owner, p.project_type, p.status,
            p.start_date, p.end_date, p.year,
            c.contract_no if c else None, cp["contract_amount"], cp["tax_rate"],
            cp["excl_tax_amount"], cp["tax_amount"], c.sign_date if c else None,
            cp["received"], cp["pending"], cp["received_ratio"],
            cp["direct_cost"], cp["indirect_cost"], cp["tax_cost"], cp["total_cost"],
            cp["profit"], cp["margin"], cp["invoiced"], cp["not_invoiced"],
            cp["recovery"]["recovery_date"],
            cp["warranty"]["amount"], cp["warranty"]["expected_date"],
            cp["warranty"]["status"], cp["next_expected_date"], p.notes,
        ]
        for col, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=_fmt(v))
            if col in money_cols:
                cell.number_format = "#,##0.00"
            elif col in date_cols and v is not None:
                cell.number_format = "yyyy-mm-dd"
        r += 1

    # 明细表
    _export_schedules(wb, projects, db, settings)
    _export_payments(wb, projects, db, settings)
    _export_costs(wb, projects, db, settings)
    _export_invoices(wb, projects, db, settings)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _iter_year_projects(db, year):
    q = db.query(Project).order_by(Project.id)
    if year:
        q = q.filter(Project.year == year)
    return q.all()


def _export_schedules(wb, projects, db, settings):
    ws = wb.create_sheet("付款计划")
    headers = ["项目名称", "节点名称", "比例(%)", "计划金额(元)", "预计到账日期", "已到账(元)", "状态", "是否质保金"]
    _style_header(ws, headers)
    _autowidth(ws, [28, 14, 10, 14, 14, 14, 10, 10])
    r = 2
    for p in projects:
        cp = compute_project(load_context(db, p), settings)
        for s in cp["schedules"]:
            for col, v in enumerate([p.name, s["name"], s["ratio_pct"], s["amount"],
                                     s["expected_date"], s["received"], s["status"],
                                     "是" if s["is_warranty"] else "否"], start=1):
                cell = ws.cell(row=r, column=col, value=_fmt(v))
                if col in (4, 6):
                    cell.number_format = "#,##0.00"
            r += 1


def _export_payments(wb, projects, db, settings):
    ws = wb.create_sheet("回款记录")
    headers = ["项目名称", "回款日期", "回款金额(元)", "类型", "对应付款节点", "备注"]
    _style_header(ws, headers)
    _autowidth(ws, [28, 12, 14, 8, 16, 24])
    r = 2
    for p in projects:
        sched_names = {s.id: s.name for s in p.schedules}
        for rec in sorted(p.payments, key=lambda x: x.payment_date):
            for col, v in enumerate([p.name, rec.payment_date, rec.amount_cents / 100.0,
                                     rec.type, sched_names.get(rec.schedule_id), rec.note], start=1):
                cell = ws.cell(row=r, column=col, value=_fmt(v))
                if col == 3:
                    cell.number_format = "#,##0.00"
            r += 1


def _export_costs(wb, projects, db, settings):
    ws = wb.create_sheet("成本明细")
    headers = ["项目名称", "类别(直接/间接)", "成本分类", "金额(元)", "发生日期", "备注"]
    _style_header(ws, headers)
    _autowidth(ws, [28, 14, 12, 14, 12, 24])
    r = 2
    for p in projects:
        for x in sorted(p.costs, key=lambda c: c.cost_date):
            kind_name = {"direct": "直接成本", "indirect": "间接成本", "tax": "税收成本"}.get(x.kind, x.kind)
            for col, v in enumerate([p.name, kind_name,
                                     x.category, x.amount_cents / 100.0, x.cost_date, x.note], start=1):
                cell = ws.cell(row=r, column=col, value=_fmt(v))
                if col == 4:
                    cell.number_format = "#,##0.00"
            r += 1


def _export_invoices(wb, projects, db, settings):
    ws = wb.create_sheet("发票记录")
    headers = ["项目名称", "发票号码", "开票金额(元)", "开票日期", "类型", "备注"]
    _style_header(ws, headers)
    _autowidth(ws, [28, 20, 14, 12, 10, 24])
    r = 2
    for p in projects:
        for inv in sorted(p.invoices, key=lambda i: i.issue_date):
            for col, v in enumerate([p.name, inv.invoice_no, inv.amount_cents / 100.0,
                                     inv.issue_date, inv.type, inv.note], start=1):
                cell = ws.cell(row=r, column=col, value=_fmt(v))
                if col == 3:
                    cell.number_format = "#,##0.00"
            r += 1


def export_projects_csv(db: Session, year: int | None = None) -> str:
    settings = get_settings(db)
    projects = _iter_year_projects(db, year)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(OVERVIEW_HEADERS)
    for p in projects:
        cp = compute_project(load_context(db, p), settings)
        c = p.contract
        writer.writerow([
            p.name, p.code, p.customer_name, p.owner, p.project_type, p.status,
            p.start_date, p.end_date, p.year,
            c.contract_no if c else None, cp["contract_amount"], cp["tax_rate"],
            cp["excl_tax_amount"], cp["tax_amount"], c.sign_date if c else None,
            cp["received"], cp["pending"], cp["received_ratio"],
            cp["direct_cost"], cp["indirect_cost"], cp["tax_cost"], cp["total_cost"],
            cp["profit"], cp["margin"], cp["invoiced"], cp["not_invoiced"],
            cp["recovery"]["recovery_date"],
            cp["warranty"]["amount"], cp["warranty"]["expected_date"],
            cp["warranty"]["status"], cp["next_expected_date"], p.notes,
        ])
    return buf.getvalue()


def export_template_xlsx() -> bytes:
    """导入模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "项目导入"
    headers = ["项目名称*", "项目编号", "客户名称", "项目负责人", "项目类型", "状态",
               "项目开始日期", "项目结束日期", "归属年度", "合同编号", "合同金额(元)",
               "税率(%)", "合同签订日期", "付款条款", "质保期(月)", "质保金比例(%)", "质保金金额(元)"]
    _style_header(ws, headers)
    _autowidth(ws, [28, 14, 20, 12, 12, 8, 12, 12, 8, 16, 14, 8, 12, 30, 10, 12, 14])
    sample = ["示例：智慧园区管理平台", "XM-2026-001", "示例客户有限公司", "张三", "软件开发", "进行中",
              "2026-03-01", "2026-12-31", 2026, "HT-2026-001", 1000000, 13, "2026-02-20",
              "合同签订后30% 进度40% 验收20% 质保金10%", 12, 10, 100000]
    for col, v in enumerate(sample, start=1):
        ws.cell(row=2, column=col, value=v)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ── 导入 ─────────────────────────────────────────────────────

COLUMN_MAP = {
    "项目名称*": "name", "项目名称": "name",
    "项目编号": "code", "客户名称": "customer_name", "项目负责人": "owner",
    "项目类型": "project_type", "状态": "status",
    "项目开始日期": "start_date", "项目结束日期": "end_date", "归属年度": "year",
    "合同编号": "contract_no", "合同金额(元)": "total_amount", "合同金额": "total_amount",
    "税率(%)": "tax_rate", "税率": "tax_rate", "合同签订日期": "sign_date",
    "付款条款": "payment_terms", "质保期(月)": "warranty_months", "质保期": "warranty_months",
    "质保金比例(%)": "warranty_ratio", "质保金金额(元)": "warranty_amount",
}


def import_projects_xlsx(db: Session, data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    headers = []
    for cell in ws[1]:
        v = cell.value
        if v is None:
            break
        headers.append(str(v).strip())
    idx = {COLUMN_MAP.get(h): i for i, h in enumerate(headers) if COLUMN_MAP.get(h)}
    if "name" not in idx:
        return {"ok": False, "error": "未找到「项目名称」列，请使用系统提供的导入模板。"}

    created, errors = 0, []
    from ..models import Contract
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        get = lambda key: row[idx[key]] if key in idx and idx[key] < len(row) else None
        name = get("name")
        if name is None or str(name).strip() == "":
            errors.append(f"第{row_no}行：项目名称为空，已跳过（示例行请删除）。")
            continue
        if str(name).startswith("示例"):
            continue
        p = Project(name=str(name).strip()[:200])
        for key, attr in [("code", "code"), ("customer_name", "customer_name"),
                          ("owner", "owner"), ("project_type", "project_type")]:
            v = get(key)
            if v is not None and str(v).strip():
                setattr(p, attr, str(v).strip())
        status = get("status")
        if status and str(status).strip() in ("未开始", "进行中", "已完成", "已结算", "已关闭"):
            p.status = str(status).strip()
        for key, attr in [("start_date", "start_date"), ("end_date", "end_date")]:
            p.__setattr__(attr, parse_date_any(get(key)))
        try:
            p.year = int(get("year")) if get("year") else None
        except (TypeError, ValueError):
            p.year = None
        db.add(p)
        db.flush()
        contract = Contract(project_id=p.id)
        contract.contract_no = str(get("contract_no")).strip() if get("contract_no") else None
        money = parse_money_text(get("total_amount"))
        if money:
            contract.total_amount_cents = yuan_to_cents(money)
        rate = get("tax_rate")
        if rate is not None and str(rate).strip() != "":
            contract.tax_rate_bp = pct_to_bp(str(rate).replace("%", ""))
        contract.sign_date = parse_date_any(get("sign_date"))
        terms = get("payment_terms")
        if terms:
            contract.payment_terms = str(terms).strip()[:4000]
        try:
            wm = get("warranty_months")
            contract.warranty_months = int(float(wm)) if wm not in (None, "") else None
        except (TypeError, ValueError):
            contract.warranty_months = None
        wr = get("warranty_ratio")
        if wr is not None and str(wr).strip() != "":
            contract.warranty_ratio_bp = pct_to_bp(str(wr).replace("%", ""))
        wa = parse_money_text(get("warranty_amount"))
        if wa:
            contract.warranty_amount_cents = yuan_to_cents(wa)
        db.add(contract)
        db.flush()
        # 程序补算税额/不含税额/质保金预计到账日
        from .calc import derive_tax
        if contract.total_amount_cents and contract.tax_rate_bp is not None:
            d_excl, d_tax = derive_tax(contract.total_amount_cents, contract.tax_rate_bp)
            contract.excl_tax_amount_cents = d_excl
            contract.tax_amount_cents = d_tax
        from .money import add_months
        if contract.warranty_months:
            start = contract.warranty_start_date or p.end_date or contract.sign_date
            if start:
                contract.warranty_expected_date = add_months(start, contract.warranty_months)
        if not p.year:
            p.year = contract.sign_date.year if contract.sign_date else (p.start_date.year if p.start_date else datetime.now().year)
        created += 1
    db.commit()
    return {"ok": True, "created": created, "errors": errors}
